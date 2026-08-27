import json
import logging
from pathlib import Path

from openpyxl import load_workbook

logger = logging.getLogger(__name__)


def ler_excel_atualizado(filepath) -> tuple[dict, dict, str, str]:
    """Lê um Excel gerado pelo Kora e retorna (resultado, dados_manuais, municipio, exercicio).

    - resultado: dict com todos os quadros extraídos dos PDFs (armazenado na aba oculta _dados_)
    - dados_manuais: dict com os campos que o usuário preencheu na aba 'Dados Manuais'
    - municipio: nome do município
    - exercicio: ano do exercício
    """
    wb = load_workbook(str(filepath), data_only=True)

    # ── 1. JSON oculto com dados extraídos ────────────────────────────────────
    resultado: dict = {}
    municipio = ""
    exercicio = ""

    if "_dados_" in wb.sheetnames:
        ws_json = wb["_dados_"]
        chunks: list[str] = []
        for row in ws_json.iter_rows(max_col=1):
            cell = row[0]
            if cell.value is not None:
                chunks.append(str(cell.value))
        json_str = "".join(chunks)
        try:
            dados = json.loads(json_str)
            resultado = dados.get("resultado", {})
            municipio = dados.get("municipio", "")
            exercicio = dados.get("exercicio", "")
        except json.JSONDecodeError as exc:
            logger.error("Falha ao parsear JSON da aba _dados_: %s", exc)
    else:
        logger.warning("Aba _dados_ não encontrada — Excel pode ter sido gerado em versão anterior")

    # ── 2. Campos manuais preenchidos pelo usuário ────────────────────────────
    dados_manuais: dict = {}

    if "Dados Manuais" in wb.sheetnames:
        ws_manual = wb["Dados Manuais"]
        for row in ws_manual.iter_rows(min_row=4, max_col=2):
            key_cell = row[0]
            val_cell = row[1] if len(row) > 1 else None

            if not (key_cell.value and isinstance(key_cell.value, str)):
                continue

            key = key_cell.value.strip()
            # Chaves de dados têm underscore; linhas de seção (cabeçalho) não
            if "_" not in key:
                continue

            if val_cell is None or val_cell.value is None:
                continue

            val = val_cell.value
            if isinstance(val, str):
                val = val.strip()
                if val == "":
                    continue
            dados_manuais[key] = val
    else:
        logger.warning("Aba 'Dados Manuais' não encontrada no Excel")

    logger.info(
        "Excel lido: %d quadros extraídos, %d campos manuais, municipio=%s, exercicio=%s",
        len(resultado), len(dados_manuais), municipio, exercicio,
    )
    return resultado, dados_manuais, municipio, exercicio
