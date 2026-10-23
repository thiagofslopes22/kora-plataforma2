"""Repasses Fundo a Fundo da Saúde (planilha do FNS — "Pesquisa Consolidada").

Fonte das duas deduções do Quadro 5 que não saem do ANX_02/ANX_10:
  - "(-) Agentes Comunitários de Saúde e de Combate às Endemias com Recursos
    Vinculados (CF, art. 198, §11)" = soma dos repasses das ações de pagamento dos
    vencimentos dos Agentes Comunitários de Saúde + Agentes de Combate às Endemias.
  - "(-) Parcela dedutível referente ao piso salarial do Enfermeiro/Técnico/Auxiliar
    de Enfermagem e Parteira (ADCT, art. 38, §2º)" = repasse da Assistência Financeira
    Complementar ao piso da enfermagem, aplicado o fator legal do exercício (2025: 91%).
"""
import logging
import re
import unicodedata

import openpyxl

logger = logging.getLogger(__name__)

# Substrings (sem acento, maiúsculas) que identificam cada ação na coluna ACAO.
_ACOES_AGENTES = (
    'AGENTES COMUNITARIOS DE SAUDE',
    'AGENTES DE COMBATE AS ENDEMIAS',
)
_ACAO_PISO_ENFERMAGEM = 'PISO SALARIAL DOS PROFISSIONAIS DA ENFERMAGEM'

# Fator legal da parcela dedutível do piso da enfermagem por exercício (ADCT art. 38 §2º).
# 2025 = 91% do total repassado. Anos sem fator definido usam 0.91 e emitem aviso no log.
_FATOR_PISO_ENFERMAGEM = {'2025': 0.91}
_FATOR_PISO_PADRAO = 0.91


def _norm(s) -> str:
    """Maiúsculas sem acento, para casar rótulos tolerando mojibake/acentuação."""
    if s is None:
        return ''
    txt = unicodedata.normalize('NFKD', str(s))
    txt = ''.join(c for c in txt if not unicodedata.combining(c))
    return txt.upper().strip()


# Abreviações comuns em nomes de município: os anexos podem trazer "GOV. LUIZ ROCHA"
# enquanto a planilha do FNS traz "GOVERNADOR LUIZ ROCHA".
_ABREVIACOES = {
    'GOV': 'GOVERNADOR', 'PRES': 'PRESIDENTE', 'STA': 'SANTA', 'STO': 'SANTO', 'S': 'SAO',
    'SR': 'SENHOR', 'DR': 'DOUTOR', 'CEL': 'CORONEL', 'MAL': 'MARECHAL', 'ENG': 'ENGENHEIRO',
    'PE': 'PADRE', 'FR': 'FREI', 'D': 'DOM', 'N': 'NOVA',
}


def _norm_municipio(s) -> str:
    tokens = re.sub(r'[^A-Z0-9 ]', ' ', _norm(s)).split()
    return ' '.join(_ABREVIACOES.get(t, t) for t in tokens)


def _num(x) -> float:
    if x is None:
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace('.', '').replace(',', '.') if ',' in str(x) else str(x).strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def _localizar_colunas(ws) -> tuple[int, dict[str, int]] | tuple[None, None]:
    """Encontra a linha de cabeçalho e mapeia as colunas de interesse (0-based)."""
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        norm = [_norm(c) for c in row]
        if 'MUNICIPIO' in norm and 'ACAO' in norm:
            cols = {}
            for idx, val in enumerate(norm):
                if val == 'MUNICIPIO':          cols['municipio'] = idx
                elif val == 'ACAO':             cols['acao'] = idx
                elif val == 'ACAO DETALHADA':   cols['acao_detalhada'] = idx
                elif val == 'VALOR TOTAL':      cols['valor_total'] = idx
                elif val == 'VALOR LIQUIDO':    cols['valor_liquido'] = idx
                elif val == 'DESCONTO':         cols['desconto'] = idx
            return i, cols
    return None, None


def extrair(xlsx_path: str, municipio: str = '', exercicio: str = '') -> dict:
    """Soma os repasses fundo a fundo das ações relevantes ao Q5.

    `municipio`: quando informado, filtra as linhas por município (comparação sem
    acento/caixa). A planilha pode conter vários municípios — sem filtro, soma todos.
    `exercicio`: define o fator legal aplicado ao piso da enfermagem (2025 → 91%).
    """
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header_row, cols = _localizar_colunas(ws)
    if not cols or 'acao' not in cols:
        wb.close()
        raise ValueError(
            "Planilha fundo a fundo sem cabeçalho reconhecível (esperado colunas "
            "'MUNICIPIO', 'ACAO', 'Valor Líquido')."
        )

    col_valor = cols.get('valor_liquido', cols.get('valor_total'))
    muni_alvo = _norm_municipio(municipio)
    if muni_alvo and 'municipio' in cols:
        nomes = {_norm_municipio(r[cols['municipio']])
                 for r in ws.iter_rows(min_row=header_row + 1, values_only=True)
                 if r and len(r) > cols['municipio'] and r[cols['municipio']]}
        if muni_alvo not in nomes and len(nomes) == 1:
            logger.warning("FF_SAUDE: município %r não aparece na planilha; usando o único município dela (%s).",
                           municipio, next(iter(nomes)))
            muni_alvo = next(iter(nomes))

    agentes = 0.0
    piso_bruto = 0.0
    n_agentes = n_piso = 0

    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if not row or len(row) <= col_valor:
            continue
        if muni_alvo and 'municipio' in cols:
            if _norm_municipio(row[cols['municipio']]) != muni_alvo:
                continue
        acao = _norm(row[cols['acao']])
        if 'acao_detalhada' in cols:
            acao = f"{acao} {_norm(row[cols['acao_detalhada']])}"
        valor = _num(row[col_valor])

        if any(termo in acao for termo in _ACOES_AGENTES):
            agentes += valor
            n_agentes += 1
        elif _ACAO_PISO_ENFERMAGEM in acao:
            piso_bruto += valor
            n_piso += 1

    wb.close()

    agentes = round(agentes, 2)
    piso_bruto = round(piso_bruto, 2)
    fator = _FATOR_PISO_ENFERMAGEM.get(exercicio, _FATOR_PISO_PADRAO)
    if exercicio not in _FATOR_PISO_ENFERMAGEM:
        logger.warning(
            "FF_SAUDE: fator do piso da enfermagem não definido para exercício %r — usando %.2f.",
            exercicio, fator,
        )
    piso_dedutivel = round(piso_bruto * fator, 2)

    logger.info(
        "FF_SAUDE: municipio=%r agentes=%.2f (%d linhas) piso_bruto=%.2f (%d linhas) "
        "fator=%.2f piso_dedutivel=%.2f",
        municipio, agentes, n_agentes, piso_bruto, n_piso, fator, piso_dedutivel,
    )
    return {
        'agentes_comunitarios_saude_endemias': agentes,
        'parcela_dedutivel_piso_enfermeiro':   piso_dedutivel,
        # Diagnóstico
        'piso_enfermagem_bruto':               piso_bruto,
        'piso_enfermagem_fator':               fator,
        'linhas_agentes':                      n_agentes,
        'linhas_piso':                         n_piso,
    }
